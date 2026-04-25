"""Genera docs/catalog_cuadros.docx con la tabla completa de cuadros + URL."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT

ROOT = Path(__file__).resolve().parents[1]  # bc_pibot/
CATALOG = ROOT / "orchestrator/catalog/catalog.json"
OUT = ROOT / "docs/catalog_cuadros.docx"


def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return str(v)


def main() -> None:
    data = json.loads(CATALOG.read_text(encoding="utf-8"))

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Cm(1.2))

    doc.add_heading("Catálogo de Cuadros — pibot", level=1)
    doc.add_paragraph(
        f"Fuente: bc_pibot/orchestrator/catalog/catalog.json · Total: {len(data)} cuadros"
    )

    headers = ["#", "Nombre del cuadro", "Indicator", "Calc mode", "Freq",
               "Price", "Seas.", "Act.", "Reg.", "Inv.", "#Series", "URL"]
    widths_cm = [0.7, 6.0, 1.5, 1.8, 0.9, 0.9, 0.9, 0.8, 1.6, 0.8, 0.9, 6.5]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.autofit = False

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.width = Cm(widths_cm[i])
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(8)

    for idx, (name, payload) in enumerate(data.items(), 1):
        cls = payload.get("classification", {}) or {}
        region_extra = f" ({cls['region']})" if cls.get("region") else ""
        values = [
            str(idx),
            name,
            _fmt(cls.get("indicator")),
            _fmt(cls.get("calc_mode")),
            _fmt(cls.get("frequency")),
            _fmt(cls.get("price")),
            _fmt(cls.get("seasonality")),
            _fmt(cls.get("has_activity")),
            _fmt(cls.get("has_region")) + region_extra,
            _fmt(cls.get("has_investment")),
            str(len(payload.get("series", []))),
            payload.get("source_url", ""),
        ]
        row = table.add_row().cells
        for i, v in enumerate(values):
            row[i].text = v
            row[i].width = Cm(widths_cm[i])
            row[i].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            for p in row[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(7.5)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
